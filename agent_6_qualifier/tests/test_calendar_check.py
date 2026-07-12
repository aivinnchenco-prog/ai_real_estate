"""Тесты проверки календарей: Google Sheets УК (заливки) и iCal."""
import sys
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import PatternFill

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "src"))

from agent8.calendar_check import (
    _blocked_days_from_sheet,
    _parse_ics_dates,
    check_ical_dates,
)

RED = PatternFill(start_color="FFFF0000", end_color="FFFF0000", fill_type="solid")
YELLOW = PatternFill(start_color="FFFFFF00", end_color="FFFFFF00", fill_type="solid")


def make_sheet():
    """Мини-копия формата УК: год в A2, March в A3, сетка чисел ниже."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "Title Test - X101"
    ws["A2"] = 2026
    ws["A3"] = "March"
    for i, d in enumerate(["M", "Tu", "W", "Th", "F", "Sa", "Su"]):
        ws.cell(row=4, column=1 + i, value=d)
    # март 2026: 1-е — воскресенье (колонка 7), дальше по сетке
    day = 1
    row, col = 5, 7
    while day <= 31:
        ws.cell(row=row, column=col, value=day)
        day += 1
        col += 1
        if col > 7:
            row, col = row + 1, 1
    # 9–15 марта красные, 24 марта жёлтый
    for r in ws.iter_rows(min_row=5, max_row=11, max_col=7):
        for c in r:
            if isinstance(c.value, int) and 9 <= c.value <= 15:
                c.fill = RED
            if c.value == 24:
                c.fill = YELLOW
    return ws


def test_sheet_red_days_detected():
    busy, tentative, covered = _blocked_days_from_sheet(make_sheet())
    assert busy == {date(2026, 3, d) for d in range(9, 16)}
    assert tentative == {date(2026, 3, 24)}
    assert (2026, 3) in covered


def test_sheet_inline_year_in_month_header():
    """Формат сводного листа: «March 2026» без отдельной ячейки-года."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "March 2026"
    ws.cell(row=3, column=1, value=5).fill = RED
    busy, _, covered = _blocked_days_from_sheet(ws)
    assert busy == {date(2026, 3, 5)}
    assert covered == {(2026, 3)}


ICS = """BEGIN:VCALENDAR
BEGIN:VEVENT
DTSTART;VALUE=DATE:20260714
DTEND;VALUE=DATE:20260718
SUMMARY:Reserved
END:VEVENT
BEGIN:VEVENT
DTSTART:20260801T140000Z
DTEND:20260805T100000Z
END:VEVENT
END:VCALENDAR"""


def test_ics_parse_events():
    events = _parse_ics_dates(ICS)
    assert (date(2026, 7, 14), date(2026, 7, 18)) in events
    assert (date(2026, 8, 1), date(2026, 8, 5)) in events


def test_ical_busy_and_free(monkeypatch):
    class R:
        text = ICS
        def raise_for_status(self):
            pass
    monkeypatch.setattr("agent8.calendar_check.requests.get", lambda *a, **k: R())

    res = check_ical_dates("https://x/cal.ics", date(2026, 7, 15), date(2026, 7, 20))
    assert res.available is False           # ночи 15-17 внутри брони 14-18

    res = check_ical_dates("https://x/cal.ics", date(2026, 7, 18), date(2026, 7, 25))
    assert res.available is True            # заезд в день чужого выезда — ок

    res = check_ical_dates("https://x/cal.ics", date(2026, 7, 20), date(2026, 7, 25))
    assert res.available is True
